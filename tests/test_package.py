import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from pdf2dify.categories import CategoryStore
from pdf2dify.config import Settings
from pdf2dify.db import Database
from pdf2dify.pipeline import PipelineRunner


def test_manual_package_has_portable_paths_and_navigation(tmp_path: Path):
    settings = Settings(
        project_root=tmp_path, data_dir=tmp_path / "data", database_path=tmp_path / "data" / "test.db",
        engine_root=tmp_path / "engine", engine_python=tmp_path / "engine" / "python.exe",
    )
    data_dir = tmp_path / "engine-data"
    export = data_dir / "full-export"
    (export / "finance").mkdir(parents=True)
    (export / "finance" / "one.docx").write_bytes(b"docx")
    custom = CategoryStore(settings.data_dir).create("供应链合规")
    (export / custom["id"]).mkdir()
    (export / custom["id"] / "two.docx").write_bytes(b"docx")
    nav = data_dir / "process-navigation" / "full-export"
    nav.mkdir(parents=True)
    (nav / "flow.docx").write_bytes(b"docx")
    content = {
        "source_count": 1, "built_sources": 1, "pending_sources": [],
        "documents": [{
            "key": "one", "domain": "finance", "source_name": "manual.pdf",
            "source_id": "source", "content_hash": "digest", "image_count": 0,
            "path": str(export / "finance" / "one.docx"), "pages": [1],
            "metadata": {"section_title": "第一章"},
        }],
    }
    content["documents"].append({
        **content["documents"][0], "key": "two", "domain": custom["id"],
        "path": str(export / custom["id"] / "two.docx"),
    })
    (export / "manifest.json").write_text(json.dumps(content), encoding="utf-8")
    nav_content = {
        "documents": [{
            "key": "nav", "domain": "process_navigation", "source_name": "manual.pdf",
            "source_id": "source", "content_hash": "digest", "image_count": 0,
            "path": str(nav / "flow.docx"), "pages": [1],
            "metadata": {"section_title": "资料导航"},
        }]
    }
    (nav / "manifest.json").write_text(json.dumps(nav_content), encoding="utf-8")
    report = data_dir / "reports" / "full-corpus-verification.json"
    report.parent.mkdir()
    report.write_text(json.dumps({"errors": [], "pending_sources": []}), encoding="utf-8")

    workspace = tmp_path / "job"
    workspace.mkdir()
    PipelineRunner(settings, Database(settings.database_path))._package(data_dir, workspace)

    manifest = json.loads((workspace / "dify-ready" / "manifest.json").read_text(encoding="utf-8"))
    assert {item["path"] for item in manifest["documents"]} == {
        "财务核算/one.docx", "供应链合规/two.docx", "资料导航/flow.docx",
    }
    assert all((workspace / "dify-ready" / item["path"]).is_file() for item in manifest["documents"])
    with zipfile.ZipFile(workspace / "dify-ready.zip") as archive:
        assert "上传说明.md" in archive.namelist()
        assert "资料导航/flow.docx" in archive.namelist()
    workbook = load_workbook(workspace / "dify-ready" / "manifest.xlsx", read_only=True)
    assert workbook.active.max_row == 4
