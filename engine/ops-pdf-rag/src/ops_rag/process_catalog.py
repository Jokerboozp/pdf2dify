"""Source-wide navigation assembled from local evidence, never invented procedures."""
import re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Inches
from .common import read_json, write_json, signature, now
from .corpus import DOMAINS, checked_asset, xml_text

VERSION = 1
IGNORED = re.compile(r'^(资料说明与适用范围|目录|版本信息|谢谢|感谢聆听|大集中ERP项目|大集中ERP.*业务组)$')


def compact(value):
    return re.sub(r'\s+', '', value or '').strip('：:')


def filename_topic(name):
    title = Path(name).stem
    title = re.sub(r'^DJZ_MDG_用户操作手册[_-]?', '', title, flags=re.I)
    title = re.sub(r'^MDG[.\d]+[\s_-]*', '', title, flags=re.I)
    title = re.sub(r'[_-]V\d+(?:\.\d+)*.*$', '', title, flags=re.I)
    title = re.sub(r'^\d+[.、\s]+', '', title)
    title = re.sub(r'\s*\(\d+\)$', '', title)
    return compact(title)


def names_from_source(source, chunks):
    """A cover field is evidence, not mandatory. Keep conflicting names visible."""
    found = []
    for chunk in chunks:
        for region in chunk['regions']:
            if region['page'] > 3:
                continue
            for table in region.get('tables', []):
                for row in table.get('rows', []):
                    if len(row) >= 2 and compact(row[0]) == '流程名称' and row[1]:
                        found.append({'name': compact(row[1]), 'page': region['page'], 'basis': 'cover_table'})
            for match in re.finditer(r'流程名称[：:\t ]*([^\n]+)', region.get('text', '')):
                candidate = compact(match.group(1))
                if 2 < len(candidate) < 100:
                    found.append({'name': candidate, 'page': region['page'], 'basis': 'cover_text'})
    unique = {f['name']: f for f in found}
    topic = filename_topic(source['name'])
    aliases = list(unique) + [topic]
    for title in list(aliases):
        short = re.sub(r'(业务流程|操作手册|用户手册|流程)$', '', title)
        if len(short) >= 3:
            aliases.append(short)
    return {'title': next(iter(unique), topic), 'aliases': list(dict.fromkeys(aliases)),
            'name_evidence': list(unique.values()), 'name_basis': 'cover' if unique else 'filename',
            'filename_topic': topic}


def outline_from_chunks(chunks):
    sections = {}
    for chunk in chunks:
        title = compact(chunk['title'])
        if not title or IGNORED.fullmatch(title) or re.fullmatch(r'[\d\W]+', title):
            continue
        key = chunk['section_index']
        item = sections.setdefault(key, {'title': chunk['title'].strip(), 'pages': [], 'chunk_ids': []})
        item['pages'] = sorted(set(item['pages']) | set(chunk['pages']))
        item['chunk_ids'].append(chunk['chunk_id'])
    return list(sections.values())


def page_label(pages):
    # Never imply missing intermediate pages are included.
    runs = []
    for p in sorted(set(pages)):
        if runs and p == runs[-1][-1] + 1:
            runs[-1].append(p)
        else:
            runs.append([p])
    return '、'.join(str(r[0]) if len(r) == 1 else f'{r[0]}–{r[-1]}' for r in runs)


def make_record(corpus):
    source, chunks = corpus['source'], corpus['chunks']
    names = names_from_source(source, chunks)
    outline = outline_from_chunks(chunks)
    domains = sorted({c['domain'] for c in chunks})
    # Prefer explicit process introductions and step tables. A large manual with
    # no such section gets a complete outline, not guessed end-to-end sequencing.
    evidence = []
    used = set()
    for c in chunks:
        if not re.search(r'业务流程简介|业务步骤详细描述|流程步骤说明', c['title']):
            continue
        for r in c['regions']:
            key = (r['page'], r.get('top'), r.get('bottom'))
            if key in used:
                continue
            used.add(key)
            text = r.get('text', '').strip()
            tables = [' | '.join((x or '').replace('\n', '') for x in row)
                      for table in r.get('tables', []) for row in table.get('rows', [])]
            if text or tables:
                evidence.append({'page': r['page'], 'text': text,
                                 'table_rows': tables, 'chunk_id': c['chunk_id']})
    has_process_description = bool(evidence)
    # Short manuals often have no formal process-summary section. Their native
    # operation text can still support a short overview (not just two headings).
    if not evidence and len(outline) <= 12:
        for c in chunks:
            if not re.search(r'操作|步骤|业务描述|流程描述', c['title']) or IGNORED.fullmatch(compact(c['title'])):
                continue
            for r in c['regions']:
                if len(r.get('text', '').strip()) < 30:
                    continue
                key = (r['page'], r.get('top'), r.get('bottom'))
                if key in used: continue
                used.add(key)
                evidence.append({'page': r['page'], 'text': r['text'], 'table_rows': [], 'chunk_id': c['chunk_id']})
    # Keep complete evidence blocks up to a bounded overview budget. Preserve an
    # explicit coverage flag rather than truncating a step/condition mid-sentence.
    kept, chars = [], 0
    for e in evidence:
        size = len(e['text']) + sum(map(len, e['table_rows']))
        if chars + size <= 10000:
            kept.append(e); chars += size
    images = []
    seen = set()
    for c in chunks:
        if not re.search(r'业务流程图', c['title']):
            continue
        for r in c['regions']:
            for im in r['images']:
                if im['path'] not in seen and len(images) < 2:
                    images.append({'page': r['page'], 'path': im['path'], 'image_id': im['image_id']})
                    seen.add(im['path'])
    return {**names, 'key': source['source_id'], 'source_id': source['source_id'],
            'source_name': source['name'], 'source_hash': source['sha256'], 'source_pages': source['pages'],
            'domains': domains, 'outline': outline, 'overview_evidence': kept,
            'evidence_complete': len(kept) == len(evidence), 'has_process_description': has_process_description,
            'images': images, 'version': VERSION}


def paragraphs(record):
    r = record
    lines = [f"流程/手册导航：{r['title']}", '可识别名称：' + '；'.join(r['aliases']),
             f"原始资料：{r['source_name']}；共{r['source_pages']}页。业务主题：" + '、'.join(DOMAINS[d] for d in r['domains']),
             '资料适用版本待确认；业务复核待完成。本卡是原始资料的导航和摘录，不是新增操作规范。']
    if r['name_evidence']:
        for e in r['name_evidence']:
            lines.append(f"原文流程名称：{e['name']}。来源：{r['source_name']}，PDF 第{e['page']}页")
        lines.append('封面流程名可能被多份分册复用；实际范围以文件主题和下列目录为准。')
    else:
        lines.append('原文未提取到明确“流程名称”字段；本卡名称取自文件名，不冒充原文正式流程名称。')
    lines.extend(['全书章节导航（顺序为原书章节顺序，不代表所有环节必须依次执行；条件分支保持独立）：'])
    for entry in r['outline']:
        lines.append(f"{entry['title']}｜来源：{r['source_name']}，PDF 第{page_label(entry['pages'])}页")
    if r['overview_evidence']:
        lines.append('原文流程说明及步骤表摘录（需要保留条件分支与原文编号）：' if r['has_process_description']
                     else '原文操作说明摘录（需要保留条件分支与原文编号；不把局部操作声称为全部流程）：')
        for e in r['overview_evidence']:
            lines.append(f"来源：{r['source_name']}，PDF 第{e['page']}页\n{e['text']}")
            if e['table_rows']:
                lines.append('原页表格逐行文本（跨页/合并关系仍需核对）：\n' + '\n'.join(e['table_rows']))
    else:
        lines.append('本卡没有抽取到独立的完整流程步骤说明。可以按上述章节给出覆盖范围/简要环节，并追问用户想展开哪一部分；不得因此声称原手册缺少正文。')
    if not r['evidence_complete']:
        lines.append('步骤说明摘录未覆盖全部说明章节；只能作为概览，详细操作需进入对应章节。')
    lines.append('用户只给流程名称时，先给有依据的概览并列出可展开部分；目录不能伪装成逐项必做的全流程。')
    return lines


def build(cfg):
    root = Path(cfg['data_dir'])
    manifest = read_json(root/'full-export/manifest.json')
    if manifest['pending_sources'] or manifest['built_sources'] != manifest['source_count']:
        raise ValueError('Process navigation requires a complete current corpus')
    sources = {d['source_id']: d for d in read_json(root/'inventory.json')['documents']}
    folder = root/'process-navigation'
    records, exports = [], []
    for sid in sorted({d['source_id'] for d in manifest['documents']}):
        corpus = read_json(root/'corpus'/f'{sid}.json')
        if corpus['source']['sha256'] != sources[sid]['sha256']:
            raise ValueError('Stale source corpus: ' + sid)
        record = make_record(corpus)
        records.append(record)
        meta = {'process_key': sid, 'source_id': sid, 'source_hash': record['source_hash'],
                'entry_kind': 'process_navigation', 'validity_status': 'unconfirmed',
                'review_status': 'pending', 'system_version': 'unknown', 'section_title': record['title']}
        dest = folder/'full-export'/f'flow-{sid}.docx'
        dest.parent.mkdir(parents=True, exist_ok=True)
        word = Document()
        for line in paragraphs(record):
            word.add_paragraph(xml_text(line))
        for im in record['images']:
            word.add_paragraph(f"原文流程配图：{record['source_name']}，PDF 第{im['page']}页；以原图核对分支与箭头。")
            word.add_picture(str(checked_asset(root, im['path'])), width=Inches(6))
        word.save(dest)
        exports.append({'key': sid, 'domain': 'process_navigation', 'source_id': sid,
                        'source_name': record['source_name'], 'path': str(dest), 'metadata': meta,
                        'image_count': len(record['images']), 'content_hash': signature(record),
                        'pages': sorted({p for e in record['outline'] for p in e['pages']})})
    write_json(folder/'catalog.json', {'created_at': now(), 'version': VERSION, 'records': records})
    write_json(folder/'full-export/manifest.json', {'created_at': now(), 'source_count': len(sources),
               'built_sources': len(records), 'pending_sources': [], 'documents': exports})
    result = {'records': len(records), 'name_basis': dict(Counter(r['name_basis'] for r in records)),
              'images': sum(len(r['images']) for r in records),
              'outline_entries': sum(len(r['outline']) for r in records)}
    write_json(root/'reports/process-navigation-build.json', result)
    return result


def selection_catalog(records):
    return '\n'.join(f"{r['key']} | {'、'.join(DOMAINS[d] for d in r['domains'])} | "
                     f"文件主题：{r['filename_topic']} | 封面/别称：{'；'.join(r['aliases'])}"
                     for r in records)
