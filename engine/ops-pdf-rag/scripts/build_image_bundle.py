"""Create a local DOCX transport bundle: each FAQ and its PNGs stay in one chunk.

The DOCX is an import container, not generated business guidance. Dify stores the
embedded images itself; no workstation HTTP server or cloud OCR is required.
"""
from pathlib import Path
import json
import hashlib
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
SEPARATOR = '<OPS_CARD_BOUNDARY>'


def plan():
    exported = json.loads((DATA / 'export/manifest.json').read_text(encoding='utf-8'))
    inventory = json.loads((DATA / 'inventory.json').read_text(encoding='utf-8'))
    docs = {d['source_id']: d for d in inventory['documents']}
    tasks = {}
    for cid in exported['included']:
        card = json.loads((DATA / 'cards' / f'{cid}.json').read_text(encoding='utf-8'))
        for source in card['sources']:
            doc = docs[source['source_id']]
            if doc['sha256'] != source['sha256']:
                raise ValueError('Card and source inventory versions disagree')
            tasks[(source['source_id'], source['page'])] = dict(source_id=doc['source_id'],
                sha256=doc['sha256'], path=doc['path'], name=doc['name'], page=source['page'],
                total_pages=doc['pages'], module=card['module'])
    target = DATA / 'image-task-manifest.json'
    target.write_text(json.dumps({'kind': 'exported_card_images', 'pages': list(tasks.values())},
                                ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Planned {len(tasks)} source pages: {target}')


def build():
    exported = json.loads((DATA / 'export/manifest.json').read_text(encoding='utf-8'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.65)
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(4)
    rows = []
    for index, cid in enumerate(exported['included']):
        card = json.loads((DATA / 'cards' / f'{cid}.json').read_text(encoding='utf-8'))
        if card['validation_issues'] or card['review_flags']:
            raise ValueError(f'Unresolved source flags: {cid}')
        text = (DATA / 'cards' / f'{cid}.md').read_text(encoding='utf-8').split('## 配图')[0].strip()
        if SEPARATOR in text:
            raise ValueError('Source contains reserved chunk delimiter')
        heading = doc.add_paragraph(card['title'], 'Title')
        heading.paragraph_format.page_break_before = index > 0
        heading.paragraph_format.keep_with_next = True
        for line in text.splitlines()[1:]:
            if line.strip():
                doc.add_paragraph(line)
        images = []
        for source in card['sources']:
            parsed = DATA / 'parsed' / source['source_id'] / f"p{source['page']:04}.json"
            if not parsed.exists():
                raise ValueError(f'Missing image extraction: {cid}')
            page = json.loads(parsed.read_text(encoding='utf-8'))
            if page.get('status') != 'complete' or page['sha256'] != source['sha256']:
                raise ValueError(f'Stale or incomplete source: {cid}')
            # Coordinate reading order is stable and independent of PDF object ordering.
            for n, im in enumerate(sorted(page['images'], key=lambda x: (x['bbox'][1], x['bbox'][0])), 1):
                path = (DATA / im['path']).resolve()
                if not path.is_relative_to((DATA / 'assets').resolve()) or not path.is_file():
                    raise ValueError('Image must be an existing project asset')
                if path.stat().st_size >= 2 * 1024 * 1024:
                    raise ValueError(f'Image exceeds Dify 2 MiB attachment limit: {path.name}')
                label = f"原文配图：{source['filename']}，PDF 第{source['page']}页，图{n}"
                caption = doc.add_paragraph(label)
                caption.paragraph_format.keep_with_next = True
                width, height = Image.open(path).size
                scale = min(7.1 / width, 5.2 / height)
                picture = doc.add_paragraph().add_run().add_picture(str(path), width=Inches(width * scale))
                picture._inline.docPr.set('descr', label)
                images.append(dict(image_id=im['image_id'], path=im['path'], caption=label,
                                   sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
        if len(images) > 10:
            raise ValueError(f'Dify supports at most 10 images per chunk: {cid}')
        if images:
            doc.add_paragraph('配图展示原文界面，图中的日期、组织编码等为示例值，不能直接作为实际输入。')
        if index < len(exported['included']) - 1:
            doc.add_paragraph(SEPARATOR)
        rows.append(dict(card_id=cid, pages=[s['page'] for s in card['sources']], images=images))
    target = DATA / 'export/ERP运维问答-图文试点.docx'
    doc.save(target)
    manifest = dict(bundle=str(target), separator=SEPARATOR, cards=len(rows),
                    cards_with_images=sum(bool(x['images']) for x in rows),
                    images=sum(len(x['images']) for x in rows), entries=rows,
                    purpose='Dify image-preserving import transport; original screenshots, no cloud OCR')
    (DATA / 'export/image-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in manifest.items() if k != 'entries'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--plan', action='store_true', help='Plan image extraction for current export')
    args = parser.parse_args()
    plan() if args.plan else build()
