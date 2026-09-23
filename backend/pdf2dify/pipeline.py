from __future__ import annotations

import json
import os
import shutil
import subprocess
import time
import zipfile
import html
import hashlib
import queue
import threading
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .config import SecretStore, Settings
from .categories import CategoryStore
from .db import Database, utcnow
from .domains import classify_path
from .jobs import JobService


class JobCancelled(Exception):
    pass


class JobPaused(Exception):
    pass


class NeedsReview(Exception):
    pass


STAGES = [
    ("scan", "扫描 PDF", 8),
    ("native", "提取原生文字和图片", 22),
    ("process", "执行 OCR", 52),
    ("build", "生成章节 DOCX", 72),
    ("navigation", "生成资料导航", 79),
    ("verify", "质量核查", 85),
    ("package", "整理 Dify 入库包", 90),
]


class PipelineRunner:
    def __init__(self, settings: Settings, db: Database):
        self.settings = settings
        self.db = db
        self.jobs = JobService(settings, db)
        self.secrets = SecretStore(settings)
        self.categories = CategoryStore(settings.data_dir)

    def run(self, job: dict) -> None:
        job_id = job["id"]
        workspace = self.settings.data_dir / "jobs" / job_id
        workspace.mkdir(parents=True, exist_ok=True)
        data_dir = workspace / "engine-data"
        data_dir.mkdir(parents=True, exist_ok=True)
        config_path = workspace / "config.yaml"
        source_root = self.jobs.source_root(job)
        self._write_engine_config(config_path, source_root, data_dir)
        self.db.update_job(job_id, output_dir=str(workspace), error=None)

        try:
            for stage, title, progress in STAGES:
                self._guard(job_id)
                if self._stage_complete(workspace, stage):
                    self.db.add_event(job_id, "info", f"跳过已完成阶段：{title}")
                    continue
                self.db.update_job(job_id, stage=stage, progress=progress, message=title)
                self.db.add_event(job_id, "info", f"开始：{title}")
                if stage == "scan":
                    self._run_engine(job_id, "scan", config_path)
                    self._load_inventory(job_id, data_dir, job.get("fixed_domain"))
                    unresolved = [item for item in self.db.list_files(job_id) if not item.get("domain")]
                    if unresolved:
                        self.db.add_event(job_id, "warning", f"有 {len(unresolved)} 个文件需要选择业务分类")
                        raise NeedsReview
                elif stage in {"native", "process", "navigation", "verify"}:
                    self._run_engine(job_id, stage, config_path)
                    if stage == "process":
                        current = self.db.get_job(job_id)
                        self.db.update_job(job_id, processed_pages=current["total_pages"])
                elif stage == "build":
                    classifications = workspace / "classifications.json"
                    mapping = {
                        item["source_id"]: item["domain"] for item in self.db.list_files(job_id)
                        if item["classification_status"] == "confirmed"
                    }
                    classifications.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
                    self._run_engine(job_id, stage, config_path, ["--classifications", str(classifications)])
                elif stage == "package":
                    self._package(data_dir, workspace)
                self._mark_stage(workspace, stage)
                self.db.add_event(job_id, "info", f"完成：{title}")

            if job["mode"] == "sync":
                self._sync_dify(job_id, config_path, workspace)
            self.db.update_job(
                job_id, status="completed", stage="completed", progress=100,
                message="处理完成", finished_at=utcnow(), processed_files=len(self.db.list_files(job_id)),
                processed_pages=self.db.get_job(job_id)["total_pages"],
            )
            self.db.set_all_file_status(job_id, "completed")
            self.db.add_event(job_id, "info", "任务已完成")
        except NeedsReview:
            self.db.update_job(job_id, status="needs_review", stage="classification", message="等待人工分类")
        except JobPaused:
            self.db.update_job(job_id, status="paused", message="任务已暂停")
            self.db.add_event(job_id, "info", "任务已暂停")
        except JobCancelled:
            self.db.update_job(job_id, status="cancelled", message="任务已取消", finished_at=utcnow())
            self.db.add_event(job_id, "warning", "任务已取消")
        except Exception as exc:
            message = f"{type(exc).__name__}: {exc}"
            self.db.update_job(job_id, status="failed", message="处理失败", error=message, finished_at=utcnow())
            self.db.add_event(job_id, "error", message)

    def _write_engine_config(self, path: Path, source_root: Path, data_dir: Path) -> None:
        config = {
            "source_root": str(source_root.resolve()),
            "data_dir": str(data_dir.resolve()),
            "ocr": {
                "engine": "rapidocr", "render_scale": 2.0, "min_score": 0.75,
                "cpu_threads": 4, "image_min_width_pt": 100, "image_min_height_pt": 45,
                "ignored_image_sha256": [],
            },
            "pilot": [],
            "categories": self.categories.mapping(),
            "dify": {"dataset_name": "pdf2dify", "indexing_technique": "high_quality", "top_k": 6,
                     "score_threshold_enabled": False},
        }
        path.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")

    def _run_engine(self, job_id: str, stage: str, config_path: Path,
                    extra: list[str] | None = None) -> None:
        if not self.settings.engine_python.is_file():
            raise FileNotFoundError(f"解析引擎 Python 不存在：{self.settings.engine_python}")
        if not (self.settings.engine_root / "src" / "ops_rag").is_dir():
            raise FileNotFoundError(f"解析引擎目录无效：{self.settings.engine_root}")
        bridge = Path(__file__).with_name("engine_bridge.py")
        command = [
            str(self.settings.engine_python), str(bridge), stage,
            "--engine-root", str(self.settings.engine_root), "--config", str(config_path),
        ] + (extra or [])
        env = os.environ.copy()
        env.update(self.secrets.read())
        if env.get("DIFY_BASE_URL"):
            env["DIFY_BASE_URL"] = env["DIFY_BASE_URL"].rstrip("/")
        process = subprocess.Popen(
            command, cwd=self.settings.engine_root, stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", env=env,
        )
        assert process.stdout is not None
        output: queue.Queue[str] = queue.Queue(maxsize=1000)
        dropped_lines = 0

        def collect_output() -> None:
            nonlocal dropped_lines
            for line in process.stdout:
                try:
                    output.put_nowait(line)
                except queue.Full:
                    dropped_lines += 1

        reader = threading.Thread(target=collect_output, daemon=True)
        reader.start()

        def drain_output(limit: int = 100) -> None:
            for _ in range(limit):
                try:
                    line = output.get_nowait()
                except queue.Empty:
                    break
                self.db.add_event(job_id, "engine", line.rstrip()[:4000])

        try:
            while True:
                drain_output()
                current = self.db.get_job(job_id)
                if current["cancel_requested"]:
                    process.terminate()
                    raise JobCancelled
                if process.poll() is not None:
                    break
                time.sleep(0.1)
            reader.join(timeout=1)
            drain_output(1000)
            if dropped_lines:
                self.db.add_event(job_id, "warning", f"引擎输出过快，省略 {dropped_lines} 行日志")
            if process.returncode:
                raise RuntimeError(f"阶段 {stage} 退出码 {process.returncode}")
        finally:
            if process.poll() is None:
                process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
            reader.join(timeout=1)

    def _load_inventory(self, job_id: str, data_dir: Path, fixed_domain: str | None) -> None:
        inventory = json.loads((data_dir / "inventory.json").read_text(encoding="utf-8"))
        files = []
        for item in inventory["documents"]:
            domain = fixed_domain or classify_path(item["relative_path"], item["name"])
            files.append({
                **item,
                "domain": domain,
                "classification_status": "confirmed" if fixed_domain else ("automatic" if domain else "needs_review"),
            })
        self.db.upsert_files(job_id, files)
        self.db.update_job(
            job_id, total_files=inventory["file_count"], total_pages=inventory["total_pages"],
            message=f"发现 {inventory['file_count']} 个 PDF，共 {inventory['total_pages']} 页",
        )
        if inventory.get("duplicates"):
            self.db.add_event(job_id, "warning", f"跳过 {len(inventory['duplicates'])} 个内容完全相同的 PDF")
        failed = [item for item in inventory["documents"] if item.get("status") != "ok"]
        if failed:
            names = "、".join(item["name"] for item in failed[:5])
            raise ValueError(f"{len(failed)} 个 PDF 无法读取：{names}")

    def _package(self, data_dir: Path, workspace: Path) -> None:
        categories = self.categories.mapping()
        source = data_dir / "full-export"
        if not (source / "manifest.json").is_file():
            raise FileNotFoundError("缺少导出清单")
        ready = workspace / "dify-ready"
        if ready.exists():
            shutil.rmtree(ready)
        ready.mkdir(parents=True)
        for domain, title in categories.items():
            domain_source = source / domain
            if domain_source.is_dir():
                shutil.copytree(domain_source, ready / f"{title}")
        navigation = data_dir / "process-navigation" / "full-export"
        if navigation.is_dir():
            (ready / "资料导航").mkdir()
            for path in navigation.glob("*.docx"):
                shutil.copy2(path, ready / "资料导航" / path.name)
        manifest = json.loads((source / "manifest.json").read_text(encoding="utf-8"))
        nav_manifest_path = data_dir / "process-navigation" / "full-export" / "manifest.json"
        nav_manifest = json.loads(nav_manifest_path.read_text(encoding="utf-8")) if nav_manifest_path.exists() else {"documents": []}
        portable = {
            "source_count": manifest.get("source_count", 0),
            "built_sources": manifest.get("built_sources", 0),
            "pending_sources": manifest.get("pending_sources", []),
            "documents": [],
        }
        for item in manifest["documents"]:
            portable["documents"].append({
                **item, "path": f"{categories[item['domain']]}/{Path(item['path']).name}"
            })
        for item in nav_manifest["documents"]:
            portable["documents"].append({
                **item, "path": f"资料导航/{Path(item['path']).name}"
            })
        (ready / "manifest.json").write_text(
            json.dumps(portable, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        self._write_manifest_xlsx(ready / "manifest.xlsx", manifest["documents"], nav_manifest["documents"], categories)
        verification_path = data_dir / "reports" / "full-corpus-verification.json"
        verification = json.loads(verification_path.read_text(encoding="utf-8"))
        self._write_report(ready / "quality-report.html", manifest, nav_manifest, verification, categories)
        self._write_upload_guide(ready / "上传说明.md", manifest, nav_manifest, categories)
        zip_path = workspace / "dify-ready.zip"
        temp = zip_path.with_suffix(".tmp")
        with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in ready.rglob("*"):
                if path.is_file():
                    archive.write(path, path.relative_to(ready).as_posix())
        os.replace(temp, zip_path)

    @staticmethod
    def _write_manifest_xlsx(path: Path, documents: list[dict], navigation: list[dict], categories: dict[str, str]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Dify入库清单"
        headers = ["类型", "业务分类", "来源文件", "章节标题", "PDF页码", "图片数", "source_id", "content_hash", "文件名"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B4C")
        for item in documents:
            sheet.append([
                "章节", categories.get(item["domain"], item["domain"]), item["source_name"],
                item.get("metadata", {}).get("section_title", ""), "、".join(map(str, item.get("pages", []))),
                item.get("image_count", 0), item["source_id"], item["content_hash"], Path(item["path"]).name,
            ])
        for item in navigation:
            sheet.append([
                "资料导航", "资料导航", item["source_name"], item.get("metadata", {}).get("section_title", ""),
                "、".join(map(str, item.get("pages", []))), item.get("image_count", 0), item["source_id"],
                item["content_hash"], Path(item["path"]).name,
            ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = [12, 16, 38, 42, 20, 10, 20, 68, 42]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + index)].width = width
        workbook.save(path)

    @staticmethod
    def _write_report(path: Path, manifest: dict, navigation: dict, verification: dict, categories: dict[str, str]) -> None:
        counts: dict[str, int] = {}
        for item in manifest["documents"]:
            counts[item["domain"]] = counts.get(item["domain"], 0) + 1
        rows = "".join(
            f"<tr><td>{html.escape(categories.get(domain, domain))}</td><td>{count}</td></tr>"
            for domain, count in sorted(counts.items())
        )
        errors = verification.get("errors", [])
        body = f"""<!doctype html><html lang='zh-CN'><meta charset='utf-8'>
<title>pdf2dify 质量报告</title><style>body{{font:15px/1.7 system-ui;max-width:980px;margin:40px auto;color:#173126}}h1{{color:#176b4c}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #d6ddd8;padding:9px;text-align:left}}th{{background:#edf4ef}}.ok{{color:#176b4c}}.bad{{color:#a9362a}}</style>
<h1>pdf2dify 技术质量报告</h1><p>来源文件：{manifest.get('source_count', 0)}；章节文档：{len(manifest['documents'])}；资料导航：{len(navigation['documents'])}。</p>
<p class='{'ok' if not errors else 'bad'}'>技术核查：{'通过' if not errors else f'发现 {len(errors)} 个错误'}。业务内容仍需业务人员复核。</p>
<h2>按知识库分类</h2><table><thead><tr><th>知识库</th><th>文档数</th></tr></thead><tbody>{rows}<tr><td>资料导航</td><td>{len(navigation['documents'])}</td></tr></tbody></table>
<h2>核查摘要</h2><pre>{html.escape(json.dumps(verification, ensure_ascii=False, indent=2))}</pre></html>"""
        path.write_text(body, encoding="utf-8")

    @staticmethod
    def _write_upload_guide(path: Path, manifest: dict, navigation: dict, categories: dict[str, str]) -> None:
        counts: dict[str, int] = {}
        for item in manifest["documents"]:
            counts[item["domain"]] = counts.get(item["domain"], 0) + 1
        rows = "\n".join(
            f"| {categories[domain]} | {count} | `{categories[domain]}/` |"
            for domain, count in sorted(counts.items())
        )
        path.write_text(
            "# Dify 手工上传说明\n\n"
            "1. 在 Dify 的「知识」页面按下表创建知识库，每个业务分类一个库。"
            "资料导航单独建库。\n"
            "2. 建议选择高质量索引、父子分段，父级按整份章节文档、子级约 500 tokens，"
            "检索选混合检索。创建前先用少量文件预览分段。\n"
            "3. 解压后进入每个分类目录，只选择其中的 `.docx` 文件，"
            "上传到表中对应的知识库。不要上传 `manifest.json` 或质量报告。\n"
            "4. 等待 Dify 索引完成，再用来源文件名、章节名和事务码做检索测试。"
            "本包的 `manifest.xlsx` 可用于核对数量和来源页码。\n\n"
            "| 知识库 | DOCX 数量 | 上传目录 |\n| --- | ---: | --- |\n"
            f"{rows}\n| 资料导航 | {len(navigation['documents'])} | `资料导航/` |\n\n"
            "文档业务有效性和截图 OCR 结果仍需业务人员复核。\n",
            encoding="utf-8",
        )

    def _sync_dify(self, job_id: str, config_path: Path, workspace: Path) -> None:
        secrets = self.secrets.read()
        if not secrets.get("DIFY_BASE_URL") or not secrets.get("DIFY_DATASET_API_KEY"):
            raise ValueError("尚未配置 Dify 地址或知识库 API Key")
        state_dir = self._sync_state_dir(secrets["DIFY_BASE_URL"])
        self._load_sync_receipts(workspace, state_dir)
        for stage, title, progress in (
            ("dify-setup", "创建或映射 Dify 知识库", 92),
            ("dify-upload", "上传并更新 Dify 文档", 95),
            ("nav-setup", "创建或映射资料导航库", 96),
            ("nav-upload", "上传资料导航", 97),
        ):
            self._guard(job_id)
            self.db.update_job(job_id, stage=stage, progress=progress, message=title)
            self.db.add_event(job_id, "info", f"开始：{title}")
            self._run_engine(job_id, stage, config_path)
            self._save_sync_receipts(workspace, state_dir)
        deadline = time.monotonic() + 12 * 60 * 60
        while True:
            self._guard(job_id)
            self.db.update_job(job_id, stage="indexing", progress=98, message="等待 Dify 索引")
            self._run_engine(job_id, "dify-status", config_path)
            self._run_engine(job_id, "nav-status", config_path)
            report = workspace / "engine-data" / "reports" / "full-indexing.json"
            state = json.loads(report.read_text(encoding="utf-8")) if report.exists() else {}
            domains = state.get("domains", {})
            pending = sum(len(value.get("target", {}).get("pending", [])) for value in domains.values())
            anchor_missing = sum(
                max(0, value.get("target", {}).get("counts", {}).get("anchor_required", 0)
                    - value.get("target", {}).get("counts", {}).get("anchors_current", 0))
                for value in domains.values()
            )
            errors = [
                entry for value in domains.values()
                for entry in value.get("target", {}).get("pending", [])
                if entry.get("status") == "error"
            ]
            nav_report = workspace / "engine-data" / "reports" / "process-navigation-status.json"
            nav_state = json.loads(nav_report.read_text(encoding="utf-8")) if nav_report.exists() else {}
            nav_pending = len(nav_state.get("pending", []))
            errors.extend(entry for entry in nav_state.get("pending", []) if entry.get("status") == "error")
            if errors:
                raise RuntimeError(f"Dify 索引有 {len(errors)} 个错误，请查看远端索引报告")
            if domains and pending == 0 and anchor_missing and nav_state and nav_pending == 0:
                self.db.update_job(job_id, stage="dify-anchor", message="补充原文检索标题")
                self._run_engine(job_id, "dify-anchor", config_path)
                self._save_sync_receipts(workspace, state_dir)
                continue
            if domains and pending == 0 and anchor_missing == 0 and nav_state and nav_pending == 0:
                break
            if time.monotonic() >= deadline:
                raise TimeoutError("等待 Dify 索引超过 12 小时，可稍后继续任务")
            time.sleep(30)
        self.db.update_job(job_id, stage="auditing", progress=99, message="核查 Dify 远端内容")
        self._run_engine(job_id, "dify-audit", config_path)
        self._run_engine(job_id, "nav-audit", config_path)
        retire_plan, source_index = self._replacement_plan(job_id, workspace, state_dir)
        if retire_plan["documents"] or retire_plan["navigation"]:
            plan_path = workspace / "retire-plan.json"
            plan_path.write_text(json.dumps(retire_plan, ensure_ascii=False, indent=2), encoding="utf-8")
            self.db.update_job(job_id, stage="retiring", message="停用已替换的旧版本")
            self._run_engine(job_id, "dify-retire", config_path, ["--retire-plan", str(plan_path)])
        self._save_sync_receipts(workspace, state_dir)
        index_path = state_dir / "source-index.json"
        temp = index_path.with_suffix(".tmp")
        temp.write_text(json.dumps(source_index, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temp, index_path)

    def _replacement_plan(self, job_id: str, workspace: Path, state_dir: Path) -> tuple[dict, dict]:
        job = self.db.get_job(job_id)
        index_path = state_dir / "source-index.json"
        index = json.loads(index_path.read_text(encoding="utf-8")) if index_path.is_file() else {}
        root = workspace / "engine-data"
        manifest = json.loads((root / "full-export/manifest.json").read_text(encoding="utf-8"))
        nav_manifest = json.loads((root / "process-navigation/full-export/manifest.json").read_text(encoding="utf-8"))
        plan = {"documents": [], "navigation": []}
        for file in self.db.list_files(job_id):
            identity = self._source_identity(job, file)
            current_keys = sorted(
                item["key"] for item in manifest["documents"]
                if item["source_id"] == file["source_id"]
            )
            current_nav = sorted(
                item["key"] for item in nav_manifest["documents"]
                if item["source_id"] == file["source_id"]
            )
            previous = index.get(identity, {})
            plan["documents"].extend(sorted(set(previous.get("documents", [])) - set(current_keys)))
            plan["navigation"].extend(sorted(set(previous.get("navigation", [])) - set(current_nav)))
            index[identity] = {
                "source_id": file["source_id"], "documents": current_keys,
                "navigation": current_nav, "updated_at": utcnow(),
            }
        return {key: sorted(set(value)) for key, value in plan.items()}, index

    @staticmethod
    def _source_identity(job: dict, file: dict) -> str:
        if job["source_type"] == "path_file":
            return Path(job["source_path"]).resolve().as_posix().lower()
        if job["source_type"] == "path_directory":
            return (Path(job["source_path"]) / file["relative_path"]).resolve().as_posix().lower()
        return "upload:" + job["name"] + ":" + file["relative_path"].replace("\\", "/").lower()

    def _sync_state_dir(self, base_url: str) -> Path:
        identity = hashlib.sha256(base_url.rstrip("/").encode("utf-8")).hexdigest()[:16]
        path = self.settings.data_dir / "dify-state" / identity
        path.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def _sync_receipt_paths(workspace: Path) -> tuple[tuple[str, str], ...]:
        return (
            ("full-state.json", "dify/full-state.json"),
            ("search-anchors.json", "dify/search-anchors.json"),
            ("navigation-state.json", "process-navigation/dify/full-state.json"),
        )

    def _load_sync_receipts(self, workspace: Path, state_dir: Path) -> None:
        root = workspace / "engine-data"
        for shared_name, relative in self._sync_receipt_paths(workspace):
            source = state_dir / shared_name
            target = root / relative
            if source.is_file() and not target.exists():
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, target)

    def _save_sync_receipts(self, workspace: Path, state_dir: Path) -> None:
        root = workspace / "engine-data"
        for shared_name, relative in self._sync_receipt_paths(workspace):
            source = root / relative
            if source.is_file():
                target = state_dir / shared_name
                temp = target.with_suffix(".tmp")
                shutil.copy2(source, temp)
                os.replace(temp, target)

    def _guard(self, job_id: str) -> None:
        current = self.db.get_job(job_id)
        if current["cancel_requested"]:
            raise JobCancelled
        if current["pause_requested"]:
            raise JobPaused

    @staticmethod
    def _marker(workspace: Path, stage: str) -> Path:
        return workspace / "checkpoints" / f"{stage}.done"

    def _stage_complete(self, workspace: Path, stage: str) -> bool:
        return self._marker(workspace, stage).is_file()

    def _mark_stage(self, workspace: Path, stage: str) -> None:
        marker = self._marker(workspace, stage)
        marker.parent.mkdir(parents=True, exist_ok=True)
        marker.write_text(utcnow(), encoding="utf-8")
