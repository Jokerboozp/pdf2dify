"""Extractive, source-bound corpus. No cloud OCR or generated business steps."""
from collections import Counter
from contextlib import closing
from pathlib import Path
import re
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from .common import read_json, write_json, signature, digest_file, now
from .pipeline import clean_lines

DOMAINS = {
    'finance': '财务核算', 'funds': '资金与预算', 'procurement': '采购与物资',
    'projects': '项目管理', 'internal': '内部交易', 'oilgas': '油气业务',
    'equipment': '设备管理', 'master_ops': '主数据操作', 'master_rules': '主数据标准',
}
SEPARATOR = '<OPS_SECTION_BOUNDARY>'
VERSION = 4
BOUNDARY_RULES_VERSION = 2


def page_number_only(text):
    return bool(re.fullmatch(r'[\d\s.,，、:：;；\-—_()/（）]+', text.strip()))


def xml_text(text):
    # Some source PDFs encode bullet glyphs as XML-forbidden controls. Keep an
    # explicit replacement marker rather than inventing the missing symbol.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '\ufffd', text)


def classify(doc, page=None):
    path, name = doc['relative_path'], doc['name']
    if name == '财务模块常见问题.pdf':
        if page in (16, 17): return 'projects'
        if page == 18: return 'equipment'
        if page == 9 or page and 50 <= page <= 55: return 'internal'
        if page and 27 <= page <= 48: return 'funds'
        return 'finance'
    if '公共数据管理标准/' in path: return 'master_rules'
    if '/主数据管理/' in path: return 'master_ops'
    if '/核算模块/' in path: return 'finance'
    if '/资金模块/' in path or '/预算模块/' in path: return 'funds'
    if '/内部交易模块/' in path: return 'internal'
    if '/油气销售/' in path or '/油气委托加工/' in path: return 'oilgas'
    if '/物资管理/' in path or '物资采购培训' in name: return 'procurement'
    if '设备管理操作手册' in name: return 'equipment'
    if any(x in path for x in ('投资项目', '其他费用项目', '数信项目', '科技项目', '费用项目操作')):
        return 'projects'
    raise ValueError('Unclassified source: ' + path)


def section_boundaries(doc, pdf, native):
    """Use PDF outline coordinates, retain role/menu/step subheadings together."""
    boundaries = [{'page': 1, 'top': 0.0, 'title': '资料说明与适用范围'}]
    if doc['name'] == '财务模块常见问题.pdf':
        for p in native:
            if p['page'] == 1 or p['page'] == 28: continue  # same question continues from p27
            match = re.search(r'问题描述[：:]([^\n]+)', p['text'])
            title = match.group(1) if match else ' / '.join(p['text'].splitlines()[1:3])
            boundaries.append({'page': p['page'], 'top': 0.0, 'title': title or doc['name']})
        return boundaries
    for item in pdf.get_toc():
        if item.page_index is None: continue
        title = re.sub(r'\s+', '', item.title)
        if re.search(r'完成该步骤|角色名称|菜单路径|详细操作描述|^\d+(?:\.\d+)*[.、]?操作步骤$', title):
            continue
        if item.level > 2: continue
        page = pdf[item.page_index]
        height = page.get_height()
        page.close()
        # Bookmark XYZ destinations use PDF bottom-origin coordinates.
        top = max(0.0, min(height, height - item.view_pos[1] - 5)) if item.view_mode == 1 and len(item.view_pos) >= 2 else 0.0
        boundaries.append({'page': item.page_index + 1, 'top': round(top, 2), 'title': item.title})
    if len(boundaries) == 1:
        # Slide decks: group consecutive pages by actual repeated section header.
        previous = None
        for p in native:
            lines = [x.strip() for x in p['text'].splitlines() if x.strip() and not page_number_only(x)]
            # Image-only continuation pages may contain only duplicated footer
            # page numbers (e.g. "80 80"). They belong to the previous procedure.
            if not lines:
                continue
            title = lines[0]
            if title != previous and p['page'] > 1:
                boundaries.append({'page': p['page'], 'top': 0.0, 'title': title[:120]})
            previous = title
    result = {}
    for b in boundaries:
        result[(b['page'], b['top'])] = b
    return sorted(result.values(), key=lambda b: (b['page'], b['top']))


def metadata(doc, domain):
    version = re.search(r'[_-](V\d+(?:\.\d+)*)', doc['name'], re.I)
    project = next((x for x in ('投资项目', '其他费用项目', '数信项目', '科技项目', '费用项目') if x in doc['relative_path']), 'not_applicable')
    return {'domain': domain, 'source_id': doc['source_id'], 'source_hash': doc['sha256'],
            'system_version': 'unknown', 'document_version': version.group(1) if version else 'unknown',
            'review_status': 'pending', 'validity_status': 'unconfirmed', 'project_type': project}


def checked_asset(root, rel):
    path = (root / rel).resolve()
    if not path.is_relative_to((root / 'assets').resolve()) or not path.is_file():
        raise ValueError('Invalid source image path')
    return path


def evidence_title(title, regions):
    """A bookmark on a repeated contents slide is navigation, not an operation."""
    def is_contents(region):
        lines = [line.strip() for line in region.get('text', '').splitlines() if line.strip()]
        return bool(lines and re.sub(r'\s+', '', lines[0]) == '目录')
    return '目录' if regions and all(is_contents(r) for r in regions) else title


def make_regions(doc, pdf, native, parsed, bounds, root):
    sections = []
    end = {'page': doc['pages'] + 1, 'top': 0.0}
    for i, start in enumerate(bounds):
        stop = bounds[i+1] if i+1 < len(bounds) else end
        regions = []
        for number in range(start['page'], min(stop['page'], doc['pages']) + 1):
            page = pdf[number-1]
            w, h = page.get_size()
            top = start['top'] if number == start['page'] else 0.0
            bottom = stop['top'] if number == stop['page'] else h
            if bottom - top < 2:
                page.close(); continue
            evidence = parsed[number]
            if top == 0 and bottom == h:
                text = evidence['cleaned_text']
            else:
                tp = page.get_textpage()
                text = tp.get_text_bounded(0, h-bottom, w, h-top)
                tp.close()
            page.close()
            text, _ = clean_lines(text)
            images = [im for im in evidence['images'] if top <= (im['bbox'][1]+im['bbox'][3])/2 < bottom]
            images = sorted(images, key=lambda im: (im['bbox'][1], im['bbox'][0]))
            # Vector-only process diagrams and tables remain visible as original page evidence.
            if not images and (('流程图' in text and len(text) < 1000) or evidence.get('tables')):
                images = [{'image_id': f"{doc['source_id']}-p{number:04}-page", 'path': evidence['page_asset'],
                           'text': '', 'lines': [], 'low_confidence_count': 0, 'bbox': [0, 0, w, h],
                           'role': 'original_page_table_or_diagram'}]
            tables = [t for t in evidence.get('tables', []) if t.get('rows') and top <= (t['bbox'][1]+t['bbox'][3])/2 < bottom]
            if text or images or tables:
                regions.append({'page': number, 'top': top, 'bottom': bottom, 'text': text, 'images': images, 'tables': tables})
        if regions:
            sections.append({'title': evidence_title(start['title'], regions), 'regions': regions,
                             'domain': classify(doc, start['page']), 'boundary_method': 'pdf_outline_or_native_heading'})
    return sections


def region_text(doc, region):
    lines = [f"来源：{doc['name']}，PDF 第{region['page']}页", region['text']]
    for table in region.get('tables', []):
        lines.append('原页表格逐行记录（保留空列；跨页及合并单元格关系以原图为准）：')
        for row in table['rows']:
            lines.append(' | '.join((cell or '').replace('\n', ' / ') for cell in row))
    for im in region['images']:
        if im.get('text'):
            lines.extend(['截图文字（本地 OCR 辅助检索；可能有识别误差，事务码以原文正文为准）：', im['text']])
    return xml_text('\n'.join(lines))


def chunk_sections(doc, sections):
    """Keep full sections where bounded; split oversized evidence with explicit continuation."""
    chunks = []
    for section_index, section in enumerate(sections):
        groups, pending, size, count = [], [], 0, 0
        for region in section['regions']:
            length = len(region_text(doc, region))
            images = len(region['images'])
            if pending and (size + length > 10000 or count + images > 10):
                groups.append(pending); pending, size, count = [], 0, 0
            # Rare dense pages are retained, with overflow images in explicit companion chunks.
            for offset in range(0, max(images, 1), 10):
                part = {**region, 'images': region['images'][offset:offset+10]}
                if offset:
                    if pending: groups.append(pending)
                    pending, size, count = [], 0, 0
                pending.append(part); size += length; count += len(part['images'])
        if pending: groups.append(pending)
        for part, regions in enumerate(groups, 1):
            cid = f"{doc['source_id']}-s{section_index+1:04}-p{part:02}"
            notes = []
            if len(groups) > 1:
                notes.append(f"本章节分为{len(groups)}部分，此为第{part}部分。回答全流程时需要其余部分依据，不得将本部分声称为全部步骤。")
            if doc['name'] == '财务模块常见问题.pdf' and any(r['page'] in (12,13,15) for r in regions):
                notes.append('原文事务码存在疑点：本条仅供核对，不得据此给出已确认的操作指令，请联系运维确认。')
            chunks.append({'chunk_id': cid, 'title': section['title'], 'domain': section['domain'],
                'section_index': section_index+1, 'part': part, 'parts': len(groups), 'notes': notes,
                'pages': sorted({r['page'] for r in regions}), 'regions': regions})
    return chunks


def build_source(root, doc):
    native = read_json(root / 'native' / (doc['source_id']+'.json'))
    if native['sha256'] != doc['sha256'] or digest_file(doc['path']) != doc['sha256']:
        raise ValueError('Source changed: ' + doc['name'])
    parsed = {}
    for n in range(1, doc['pages'] + 1):
        path = root / 'parsed' / doc['source_id'] / f'p{n:04}.json'
        p = read_json(path) if path.exists() else {}
        if p.get('status') != 'complete' or p.get('sha256') != doc['sha256']:
            return None
        parsed[n] = p
    with closing(pdfium.PdfDocument(doc['path'])) as pdf:
        bounds = section_boundaries(doc, pdf, native['pages_data'])
        sections = make_regions(doc, pdf, native['pages_data'], parsed, bounds, root)
    chunks = chunk_sections(doc, sections)
    outputs = []
    for domain in sorted({x['domain'] for x in chunks}):
        selected = [x for x in chunks if x['domain'] == domain]
        # One procedure/section per full-doc parent preserves context beyond the
        # server's default 4,000-token paragraph-parent limit.
        bundles = [[chunk] for chunk in selected]
        for bi, bundle in enumerate(bundles, 1):
            key = f"{doc['source_id']}-{domain}-{bi:03}"
            filename = f"ops-{key}.docx"
            dest = root / 'full-export' / domain / filename
            dest.parent.mkdir(parents=True, exist_ok=True)
            word = Document()
            for ci, chunk in enumerate(bundle):
                if ci: word.add_paragraph(SEPARATOR)
                word.add_heading(xml_text(f"{doc['name']} / {chunk['title']}"), level=1)
                word.add_paragraph(f"业务主题：{DOMAINS[domain]}。资料适用版本待确认；业务复核状态：待复核。")
                for note in chunk['notes']: word.add_paragraph(note)
                for region in chunk['regions']:
                    word.add_paragraph(region_text(doc, region))
                    for j, im in enumerate(region['images'], 1):
                        word.add_paragraph(f"原文配图：{doc['name']}，PDF 第{region['page']}页，图{j}；截图中的数值为示例。")
                        word.add_picture(str(checked_asset(root, im['path'])), width=Inches(6))
            word.save(dest)
            outputs.append({'key': key, 'path': str(dest), 'domain': domain, 'source_id': doc['source_id'],
                'source_name': doc['name'], 'metadata': {**metadata(doc, domain), 'section_title': bundle[0]['title']},
                'chunk_ids': [x['chunk_id'] for x in bundle], 'pages': sorted({p for x in bundle for p in x['pages']}),
                'image_count': sum(len(r['images']) for x in bundle for r in x['regions']),
                'content_hash': signature({'version': VERSION, 'metadata': metadata(doc, domain), 'chunks': bundle})})
    result = {'source': doc, 'chunks': chunks, 'outputs': outputs, 'version': VERSION,
              'boundary_rules_version': BOUNDARY_RULES_VERSION, 'created_at': now()}
    write_json(root / 'corpus' / f"{doc['source_id']}.json", result)
    return result


def build(cfg, limit=None):
    root = Path(cfg['data_dir'])
    docs = read_json(root / 'inventory.json')['documents']
    results, pending = [], []
    done = 0
    for doc in docs:
        cached_path = root / 'corpus' / (doc['source_id']+'.json')
        cached = read_json(cached_path) if cached_path.exists() else {}
        boundary_upgrade = cached.get('boundary_rules_version', 1) < BOUNDARY_RULES_VERSION and any(
            page_number_only(c['title']) for c in cached.get('chunks', []))
        if not boundary_upgrade and cached.get('version') == VERSION and cached.get('source', {}).get('sha256') == doc['sha256'] and all(Path(x['path']).exists() for x in cached.get('outputs', [])):
            results.append(cached); continue
        if limit is not None and done >= limit:
            pending.append(doc['source_id']); continue
        result = build_source(root, doc)
        if result:
            results.append(result); done += 1
            print(f"Built {doc['name']}: {len(result['chunks'])} parent chunks", flush=True)
        else: pending.append(doc['source_id'])
    outputs = [x for r in results for x in r['outputs']]
    report = {'created_at': now(), 'source_count': len(docs), 'built_sources': len(results),
              'processed_pages': sum(r['source']['pages'] for r in results), 'pending_sources': pending,
              'parent_chunks': sum(len(r['chunks']) for r in results),
              'image_references': sum(x['image_count'] for x in outputs),
              'domains': dict(Counter(x['domain'] for x in outputs)), 'documents': outputs}
    write_json(root / 'full-export' / 'manifest.json', report)
    return {k:v for k,v in report.items() if k != 'documents'}
